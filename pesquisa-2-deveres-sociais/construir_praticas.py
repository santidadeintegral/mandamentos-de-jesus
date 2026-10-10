#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Pesquisa 2, etapa 2 — Lista de práticas do voluntário (checklist à maneira metodista). Constrói em docs/02-lista-de-praticas/:

  Lista_de_Praticas_do_Voluntario.docx/.pdf   — checklist completo: ação, como fazer agora, base bíblica, marcação hoje/semana/mês
  Cartao_de_Bolso.docx/.pdf                   — uma página, 52 práticas com caixa de marcação
  Lista_de_Praticas_Rastreabilidade.xlsx      — cada prática ligada aos versículos julgados na etapa 1, com justificativa

Uso: python3 pesquisa-2-deveres-sociais/construir_praticas.py
"""
import collections
import datetime
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent
sys.path.insert(0, str(REPO))
sys.path.insert(0, str(REPO / "comum"))
sys.path.insert(0, str(ROOT / "dados"))
from timbrado import *  # noqa: E402,F401
from biblia import jesus_verses, ref_key, ref_pt  # noqa: E402
from rotulos import ROTULOS  # noqa: E402
from categorias import CATEGORIAS  # noqa: E402
from praticas import GRUPOS, PRATICAS  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

ROOT = Path(__file__).resolve().parent  # reafirma: o star-import de timbrado traz outro ROOT
DOCS = ROOT / "docs" / "02-lista-de-praticas"
DOCS.mkdir(parents=True, exist_ok=True)
LIGHT_GREY = RGBColor(0x7A, 0x7A, 0x7A)
TIPO_PT = {"O": "ordem direta", "E": "ensino normativo", "P": "parábola/ilustração", "J": "juízo escatológico",
           "B": "bem-aventurança", "A": "advertência/'ai'", "M": "modelo de Jesus", "C": "citação da Escritura"}
FORCA_PT = {"D": "direta", "I": "indireta", "S": "situacional"}
GRUPO_NOME = {g: n for g, n, _ in GRUPOS}
GRUPO_LEMA = {g: l for g, _, l in GRUPOS}


# ----------------------------------------------------------------- dados
def load():
    textos = {r: wj for r, wj, _ in jesus_verses()}
    usados = collections.defaultdict(list)
    for p in PRATICAS:
        for r in [p["base"]] + p["refs"]:
            if r not in textos:
                raise SystemExit(f"prática {p['id']}: versículo fora do corpus: {r}")
            if r not in ROTULOS:
                raise SystemExit(f"prática {p['id']}: versículo não julgado na etapa 1: {r}")
            usados[r].append(p["id"])
    nao_usados = [r for r in sorted(ROTULOS, key=ref_key) if r not in usados]
    return textos, usados, nao_usados


def justificativa(p):
    cats, tipo, forca, dest, det, resumo, *nota = ROTULOS[p["base"]]
    alvo = {"A": "uma pessoa com carência corporal ao alcance (fome, sede, frio, doença, prisão, ferimento)",
            "B": "uma pessoa ao alcance que precisa ouvir, ser buscada ou fortalecida",
            "C": "uma pessoa com quem há ofensa, dívida ou conflito pendente",
            "D": "uma pessoa que pede, precisa ou não pode retribuir",
            "E": "os pais, o cônjuge ou um dependente",
            "F": "quem está ao redor no trabalho, na casa ou na igreja",
            "G": "o trabalhador, o credor, o Estado, o vizinho ou o explorado"}[p["grupo"]]
    return (f"Ação concreta que pode ser iniciada de imediato, pressupondo {alvo}. Base: {ref_pt(p['base'])} — {resumo} "
            f"({TIPO_PT[tipo]}, força {FORCA_PT[forca]}{', com detalhe prático de Jesus' if det else ''}; categoria(s) "
            f"{', '.join(CATEGORIAS[c]['titulo'] for c in cats.split(';'))}). Sustentada por mais {len(p['refs'])} passagem(ns) julgada(s) na etapa 1."
            + (f" Paralelo metodista: {p['wesley']}." if p["wesley"] else ""))


# ----------------------------------------------------------------- XLSX
def xlsx(textos, usados, nao_usados):
    wb = Workbook()
    H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="2D5130")
    WRAP = Alignment(wrap_text=True, vertical="top")

    def sheet(title, headers, data, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c in ws[1]:
            c.font = H; c.fill = HF; c.alignment = WRAP
        for r in data:
            ws.append(list(r))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    leia = wb.create_sheet("Leia-me")
    leia.column_dimensions["A"].width = 120
    for line in [
        "LISTA DE PRÁTICAS DO VOLUNTÁRIO — RASTREABILIDADE (ONG Santidade Integral em Cristo)",
        "",
        "Etapa 2 da pesquisa 'Deveres sociais ensinados por Jesus'. A etapa 1 (docs/01-levantamento-inicial) leu os 1.955 versículos com fala de Jesus e julgou 384.",
        "Esta etapa converte esses versículos em ações concretas que uma pessoa pode parar o que está fazendo e realizar imediatamente, "
        "pressupondo que todo tipo de pessoa está disponível ao redor (com fome, doente, presa, só, em dívida…).",
        "Inspiração: 'Regras Gerais das Sociedades Unidas' (John Wesley, 1743) — fazer o bem aos corpos e às almas — e o exame diário do Holy Club (Oxford, 1729).",
        "",
        "Abas:",
        "1. Práticas — as 52 práticas: grupo, ação, como fazer agora, versículo-base, demais versículos, paralelo metodista, justificativa.",
        "2. Práticas x Versículos — uma linha por par (prática, versículo), com o julgamento da etapa 1 (categoria, tipo, força, resumo) e o texto.",
        "3. Versículos não usados — versículos julgados na etapa 1 que não geraram prática imediata, com o motivo.",
        "",
        "Gerado por construir_praticas.py em " + datetime.date.today().isoformat() + ".",
    ]:
        leia.append([line])

    data = [(p["id"], p["grupo"], GRUPO_NOME[p["grupo"]], p["acao"], p["como"], ref_pt(p["base"]),
             "; ".join(ref_pt(r) for r in p["refs"]), 1 + len(p["refs"]), p["wesley"], justificativa(p)) for p in PRATICAS]
    sheet("Práticas", ["#", "Grupo", "Nome do grupo", "Ação", "Como fazer agora", "Versículo-base", "Demais versículos", "Nº versículos",
                       "Paralelo metodista", "Justificativa (por que entra na lista)"], data, [5, 6, 28, 36, 60, 12, 45, 8, 40, 90])

    data = []
    for p in PRATICAS:
        for i, r in enumerate([p["base"]] + p["refs"]):
            cats, tipo, forca, dest, det, resumo, *nota = ROTULOS[r]
            data.append((p["id"], p["acao"], ref_pt(r), "base" if i == 0 else "apoio", cats, TIPO_PT[tipo], FORCA_PT[forca],
                         "sim" if det else "", resumo, textos[r]))
    sheet("Práticas x Versículos", ["#", "Ação", "Versículo", "Papel", "Categoria(s)", "Tipo", "Força", "Detalhe prático", "Resumo (etapa 1)", "Texto (WEB)"],
          data, [5, 36, 10, 7, 14, 16, 10, 8, 45, 70])

    data = []
    for r in nao_usados:
        cats, tipo, forca, dest, det, resumo, *nota = ROTULOS[r]
        motivo = {"S": "ordem situacional (presa à ocasião histórica)",
                  "P": "versículo narrativo de parábola; a prática derivada usa o versículo-chave da mesma parábola",
                  "J": "versículo de enquadramento do juízo das nações; as obras estão em Mt 25:35-36",
                  "A": "advertência contra opositores; o dever positivo correspondente já está em outra prática",
                  "M": "exemplo de Jesus que fundamenta uma prática já listada",
                  "E": "ensino que explica ou motiva uma prática já listada, sem ação adicional",
                  "C": "citação que fundamenta uma prática já listada",
                  "O": "ordem já coberta por prática equivalente ou dirigida a situação específica",
                  "B": "bem-aventurança: descreve o caráter, não uma ação imediata"}
        data.append((ref_pt(r), cats, TIPO_PT[tipo], FORCA_PT[forca], resumo, motivo[tipo] if forca != "S" else motivo["S"]))
    sheet("Versículos não usados", ["Versículo", "Categoria(s)", "Tipo", "Força", "Resumo (etapa 1)", "Motivo de não gerar prática própria"],
          data, [10, 14, 16, 10, 50, 70])
    out = DOCS / "Lista_de_Praticas_Rastreabilidade.xlsx"
    wb.save(out)
    print("OK ->", out)


# ----------------------------------------------------------------- DOCX helpers
def finish(doc, name):
    out = DOCS / f"{name}.docx"
    doc.save(out)
    print("OK ->", out)
    print("OK ->", to_pdf(out))


def checklist_table(doc, praticas):
    widths = [0.8, 7.5, 3.2, 4.1, 1.8]
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    set_cell_margins(t, top=30, bottom=30)
    for i, h in enumerate(["#", "Prática — o que fazer agora", "Base bíblica", "Fiz?  hoje / semana / mês", "Última vez"]):
        c = t.rows[0].cells[i]
        shade(c, GREEN_HEX)
        p = c.paragraphs[0]; spacing(p, after=0)
        run(p, h, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
    for pr in praticas:
        cells = t.add_row().cells
        for c in cells:
            spacing(c.paragraphs[0], after=0)
        run(cells[0].paragraphs[0], str(pr["id"]), bold=True, size=9, color=GREEN)
        p = cells[1].paragraphs[0]
        run(p, pr["acao"], bold=True, size=9.5)
        p2 = cells[1].add_paragraph(); spacing(p2, after=0)
        run(p2, pr["como"], size=8, color=GREY)
        p = cells[2].paragraphs[0]
        run(p, ref_pt(pr["base"]), bold=True, size=8.5, color=GREEN)
        if pr["refs"]:
            p2 = cells[2].add_paragraph(); spacing(p2, after=0)
            run(p2, "; ".join(ref_pt(r) for r in pr["refs"]), italic=True, size=7, color=LIGHT_GREY)
        p = cells[3].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, f"{BOX} hoje   {BOX} sem.   {BOX} mês", size=9)
        p = cells[4].paragraphs[0]
        run(p, "___/___", size=9, color=GREY)
        cant_split(t.rows[-1])
    fix_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ----------------------------------------------------------------- DOCX 1: checklist completo
def doc_lista(usados):
    doc = new_document(
        header_meta="Voluntário(a): ______________________________    Mês/ano: ____________    Grupo/célula: ______________",
        footer_label="Lista de Práticas do Voluntário — 52 ações que Jesus mandou fazer ao próximo  •  v1.0",
        top_margin_cm=3.8, bottom_margin_cm=1.4,
    )
    title_block(doc, "LISTA DE PRÁTICAS DO VOLUNTÁRIO",
                "O que Jesus mandou fazer ao próximo, em ações que você pode realizar agora",
                "Vá e faça o mesmo.", "Lucas 10:37", after=4)
    section_title(doc, "", "Como usar esta lista")
    body(doc, f"Esta lista reúne {len(PRATICAS)} ações concretas tiradas das falas de Jesus nos quatro Evangelhos (etapa 1 deste estudo: {len(ROTULOS)} versículos lidos e julgados). "
              "Cada linha é algo que uma pessoa pode parar o que está fazendo e realizar imediatamente, porque sempre há alguém ao alcance: "
              "com fome, com sede, doente, preso, de fora, em dívida, ofendido, sozinho.", size=9.5)
    body(doc, "O método é o dos primeiros metodistas. As 'Regras Gerais' de John Wesley (1743) pediam que cada membro fizesse o bem 'de todas as formas possíveis, "
              "a todas as pessoas': aos corpos, dando de comer ao faminto, vestindo o nu, visitando os doentes e os presos; às almas, instruindo, repreendendo "
              "e exortando. No Holy Club de Oxford, John e Charles Wesley e George Whitefield examinavam-se todos os dias com perguntas do tipo "
              "'visitei um doente ou um preso hoje?'. George Whitefield fundou um orfanato; Wesley abriu dispensário, escola e fundo de empréstimo para os pobres.", size=9.5)
    bullet(doc, "Leia a lista no início do dia ou da semana e escolha ao menos uma prática para realizar.", lead="1.", size=9)
    bullet(doc, "Ao fazer, marque 'hoje', 'semana' ou 'mês' e anote a data na coluna 'Última vez'.", lead="2.", size=9)
    bullet(doc, "No encontro do grupo, compartilhe o que fez e o que ficou por fazer. Nada aqui é mérito: 'somos servos inúteis; fizemos o que devíamos' (Lc 17:10).", lead="3.", size=9)
    bullet(doc, "Em cada linha, o versículo em verde é a base da prática; os demais, em cinza, são as outras passagens em que Jesus ensina o mesmo.", lead="4.", size=9)
    note(doc, "Abreviaturas: Mt = Mateus, Mc = Marcos, Lc = Lucas, Jo = João. A planilha de rastreabilidade mostra, para cada prática, o julgamento de cada versículo.", size=8)

    for g, nome, lema in GRUPOS:
        itens = [p for p in PRATICAS if p["grupo"] == g]
        section_title(doc, g + ".", f"{nome}  ({len(itens)} práticas)")
        note(doc, lema, size=8.5)
        checklist_table(doc, itens)

    doc.add_page_break()
    section_title(doc, "", "Exame semanal (à maneira do Holy Club)")
    for q in ["A quem dei de comer, de beber ou de vestir esta semana?",
              "Quem visitei: doente, preso, sozinho, afastado?",
              "A quem perdoei, com quem me reconciliei, a quem repreendi a sós?",
              "O que dei em secreto, sem esperar retorno? A quem emprestei?",
              "A quem falei de Jesus ou convidei? A quem ensinei ou fortaleci?",
              "Fui justo com quem trabalhou para mim, com meus pais, com o Estado, com o adversário?",
              "Que prática desta lista nunca fiz? O que farei amanhã?"]:
        bullet(doc, q, lead=BOX, size=9.5, after=2)
    p = doc.add_paragraph(); spacing(p, before=8, after=0)
    run(p, "Resumo do mês — práticas realizadas: ______ de 52      Assinatura do voluntário: ______________________________", size=9)
    finish(doc, "Lista_de_Praticas_do_Voluntario")


# ----------------------------------------------------------------- DOCX 2: cartão de bolso (uma página, duas colunas)
def doc_cartao():
    doc = new_document(header_meta="Cartão de bolso do voluntário    •    Nome: ______________________    Mês: ________",
                       footer_label="Cartão de Bolso — 52 práticas  •  v1.0", top_margin_cm=3.6, bottom_margin_cm=1.2)
    title_block(doc, "52 PRÁTICAS PARA FAZER HOJE", "Marque ao realizar. Sempre há alguém ao alcance.",
                "Vá e faça o mesmo.", "Lucas 10:37", after=4)
    # duas colunas a partir daqui
    new_sec = doc.add_section()
    new_sec.left_margin = new_sec.right_margin = Cm(1.5)
    new_sec.top_margin = Cm(3.6); new_sec.bottom_margin = Cm(1.2)
    sect_pr = new_sec._sectPr
    cols = OxmlElement("w:cols"); cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "500"); sect_pr.append(cols)
    typ = OxmlElement("w:type"); typ.set(qn("w:val"), "continuous"); sect_pr.insert(0, typ)
    for g, nome, _ in GRUPOS:
        p = doc.add_paragraph(); spacing(p, before=4, after=1); p.paragraph_format.keep_with_next = True
        run(p, f"{g}. {nome}", bold=True, size=8.5, color=GREEN)
        for pr in [x for x in PRATICAS if x["grupo"] == g]:
            p = doc.add_paragraph(); spacing(p, after=0)
            p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.45)
            run(p, f"{BOX} ", size=9)
            run(p, f"{pr['id']}. {pr['acao']} ", size=8)
            run(p, f"({ref_pt(pr['base'])})", italic=True, size=7, color=LIGHT_GREY)
    # seção contínua final: faz o LibreOffice/Word equilibrarem as duas colunas
    end_sec = doc.add_section()
    end_pr = end_sec._sectPr
    typ = OxmlElement("w:type"); typ.set(qn("w:val"), "continuous"); end_pr.insert(0, typ)
    finish(doc, "Cartao_de_Bolso")


if __name__ == "__main__":
    textos, usados, nao_usados = load()
    print(f"práticas: {len(PRATICAS)} | versículos usados: {len(usados)} de {len(ROTULOS)} | não usados: {len(nao_usados)}")
    xlsx(textos, usados, nao_usados)
    doc_lista(usados)
    doc_cartao()
